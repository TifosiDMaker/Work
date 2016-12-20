from docx import Document

dest_path = 'd:/Tifosi/12月/OT5865/OT5865-泰安特种车-中译英-手册、说明书/External Review/en-US/零部件图册/00组总图.doc.review.docx'
doc = Document(dest_path)
t = doc.tables[0]

for j in range(len(t.rows)):
    if t.cell(j, 2).text == "":
        continue
    print(t.cell(j, 2).text)