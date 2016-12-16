import docx
from docx import Document
from docx import table
from docx.table import Table


doc = Document('d:/Tifosi/12月/WM318/WM318-长客-中译英-技术材料（包括图纸）/External Review/en-US/16-567-1_CDRL_18-18_Corrosion_Control_Plan-=B20161208_译红字,对照.docx.review.docx')
t = doc.tables[0]
for i in range(len(t.rows)):
    print(t.cell(i,2).text)
print('循环完了')


#s = t.columns.table
#print (s)