import win32com.client
from docx import Document

word = win32com.client.gencache.EnsureDispatch('Word.Application')
thisdoc = word.Documents.Open(FileName = 'd:\\Tifosi\\12月\\OT5865\\OT5865-泰安特种车-中译英-手册、说明书\\External Review\\en-US\\零部件图册\\00组总图.doc.review.docx')
word.Visible = 0

dest_path = 'd:/Tifosi/12月/OT5865/OT5865-泰安特种车-中译英-手册、说明书/External Review/en-US/零部件图册/00组总图.doc.review.docx'
doc = Document(dest_path)
t = doc.tables[0]

for rcount in range(thisdoc.Tables(1).Rows.Count):
    if thisdoc.Tables(1).Cell(int(rcount), 1).Shading.BackgroundPatternColorIndex != 8:
        #thisdoc.Tables(1).Rows(int(rcount)).Delete
        thisdoc.Tables(1).Cell(int(rcount), 3).Range.Text = ""
        rcount -= 1
thisdoc.Save()
word.ActiveDocument.Close()
#for j in range(len(t.rows)):
#    if t.cell(j, 2).text == "":
#        continue
#    print(t.cell(j, 2).text)
#word.DisplayAlert = False
